# -*- coding: utf-8 -*-
"""
generar_manual_estudiante.py
Generador del Manual del Estudiante (docx) para I.E.P. Amancio Varona.

Usa los mismos estilos que el Manual del Auxiliar para que todos los manuales
se vean iguales. Todo lo que dice está comprobado contra el campus: si cambia
una pantalla, cambia aquí y vuelve a generar.

    .venv\\Scripts\\python.exe scripts\\generar_manual_estudiante.py
"""

import os
import sys

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from generar_manual_auxiliar import (  # noqa: E402
    COLOR_AZUL, COLOR_GRIS_CLARO, COLOR_GRIS_MEDIO, COLOR_GRIS_OSCURO, COLOR_GUINDA,
    add_bullet, add_callout, add_captura_box, add_custom_table, add_heading_1,
    add_heading_2, add_heading_3, add_p, crear_documento,
)

DOC_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "Manuales", "Manual_del_Estudiante_Campus_Virtual.docx",
)


def pasos(doc, lista):
    for i, texto in enumerate(lista, 1):
        add_bullet(doc, f"{i}. ", texto)


def puntos(doc, lista):
    for negrita, texto in lista:
        add_bullet(doc, f"• {negrita}" if negrita else "• ", texto)


def construir_manual():
    doc = crear_documento()

    # ------------------------------------------------------------- PORTADA
    add_p(doc, "I.E.P. AMANCIO VARONA — TUMÁN", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, color=COLOR_AZUL, space_after=2)
    add_p(doc, "Campus Virtual", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=11, color=COLOR_GRIS_MEDIO, space_after=6)
    add_p(doc, "MANUAL DEL ESTUDIANTE", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=24, bold=True, color=COLOR_GUINDA, space_after=4)
    add_p(doc, "Todo lo que necesitas saber para usar tu campus virtual, explicado paso a paso", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13, color=COLOR_GRIS_OSCURO, space_after=4)
    add_p(doc, "Estudiante · Año escolar 2026", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=11, color=COLOR_GRIS_CLARO, space_after=14)

    add_callout(
        doc,
        "Lo más importante, en tres líneas",
        "1) Tu usuario es ALU- seguido de tu DNI (por ejemplo, ALU-73715514).   "
        "2) La primera vez, tu contraseña es tu número de DNI.   "
        "3) Apenas entres, el campus te pedirá cambiarla: elige una contraseña nueva y segura, "
        "que solo tú conozcas. Mientras sigas usando tu DNI, cualquiera que lo sepa podría entrar a tu cuenta.",
        theme="guinda",
    )

    add_callout(
        doc,
        "Cómo usar este manual",
        "No hace falta leerlo de corrido. Si es tu primera vez, empieza por la Parte 1. "
        "Después busca solo lo que necesites: cursos y tareas en la Parte 3, notas y asistencia en la Parte 4, "
        "pagos y trámites en la Parte 5. Al final hay preguntas frecuentes.",
        theme="azul",
    )

    # ------------------------------------------------------ PARTE 1. INGRESO
    add_heading_1(doc, "Parte 1. Tu primer ingreso")

    add_heading_2(doc, "1.1 Tus datos de acceso")
    add_p(doc, "Para entrar al campus necesitas dos cosas: tu usuario y tu contraseña.")
    add_custom_table(
        doc,
        ["Dato", "Qué escribir", "Ejemplo (DNI 73715514)"],
        [
            ["Usuario", "Las letras ALU, un guion y tu DNI. Todo junto, sin espacios.", "ALU-73715514"],
            ["Contraseña (primera vez)", "Tu número de DNI, las 8 cifras.", "73715514"],
            ["Contraseña (después)", "La contraseña nueva que tú elijas en tu primer ingreso.", "Solo la sabes tú"],
        ],
        [Inches(1.7), Inches(3.2), Inches(1.6)],
    )
    add_callout(
        doc,
        "Ojo con el usuario: no es solo tu DNI",
        "Aunque la casilla dice «Usuario / DNI», si escribes solo los números el campus no te dejará entrar. "
        "Tienes que poner delante ALU y un guion: ALU-73715514. Es el error más común.",
        theme="azul",
    )

    add_heading_2(doc, "1.2 Cómo entrar al campus")
    pasos(doc, [
        "Abre tu navegador (Chrome, Edge o Firefox), en la computadora o en el celular, y entra a la página del colegio.",
        "Haz clic en el botón CAMPUS VIRTUAL.",
        "En «Usuario / DNI» escribe tu usuario completo: ALU- y tu DNI.",
        "En «Contraseña» escribe tu DNI (si es la primera vez) o tu contraseña nueva.",
        "Presiona «Iniciar Sesión».",
    ])
    add_captura_box(doc, 1, "Pantalla de acceso al campus con los campos de usuario y contraseña.")

    add_heading_2(doc, "1.3 La primera vez: cambia tu contraseña")
    add_p(doc, "La primera vez que entres, el campus te llevará a la pantalla «Completa tu primer ingreso». "
               "No podrás usar el campus hasta completarla, y no tiene botón para salir: es a propósito, para proteger tu cuenta.")
    add_p(doc, "Te pedirá lo siguiente:")
    puntos(doc, [
        ("Correo de tus padres o apoderado: ", "escríbelo dos veces (en la segunda casilla no se puede pegar, hay que escribirlo). "
         "Ahí llegarán los avisos de asistencia, notas y trámites, así que debe ser un correo que tus padres revisen."),
        ("Contraseña actual: ", "tu número de DNI."),
        ("Nueva contraseña: ", "la que tú elijas. Debe tener al menos 8 caracteres y no puede ser tu DNI."),
        ("Confirmar nueva contraseña: ", "escribe la misma otra vez."),
    ])
    add_p(doc, "Presiona «Guardar y continuar». Si todo está bien, verás un mensaje verde y el campus te llevará a tu página de inicio.")
    add_captura_box(doc, 2, "Pantalla «Completa tu primer ingreso» con el correo y el cambio de contraseña.")

    add_callout(
        doc,
        "¿Por qué hay que cambiar la contraseña?",
        "Tu DNI no es un secreto: lo conocen tus compañeros, aparece en listas y documentos. "
        "Si sigues entrando con tu DNI, cualquiera podría abrir tu cuenta, ver tus notas, leer tus mensajes "
        "o pedir trámites a tu nombre. Con una contraseña propia, solo entras tú.",
        theme="guinda",
    )

    add_heading_2(doc, "1.4 Cómo crear una contraseña segura")
    add_p(doc, "Mientras escribes la contraseña nueva, una barra de color te dice qué tan segura es: "
               "Débil (roja), Media (amarilla) o Fuerte (verde). Intenta llegar siempre a Fuerte.")
    add_custom_table(
        doc,
        ["Haz esto", "Evita esto"],
        [
            ["Usa 10 caracteres o más.", "Tu DNI, tu fecha de nacimiento o tu número de celular."],
            ["Mezcla letras mayúsculas, minúsculas y números.", "Tu nombre, el de tu mascota o el del colegio."],
            ["Une palabras que solo tú relacionas, por ejemplo una frase corta con números.", "12345678, abcdefgh, contraseña o qwerty."],
            ["Elige algo que puedas recordar sin tener que anotarlo.", "La misma contraseña que usas en redes sociales o juegos."],
        ],
        [Inches(3.25), Inches(3.25)],
    )
    add_callout(
        doc,
        "Tu contraseña es solo tuya",
        "No se la des a nadie: ni a tus amigos, ni a tus compañeros. Nadie del colegio te la va a pedir nunca, "
        "ni por mensaje, ni por teléfono, ni en persona. Si alguien la pide, no la des y avisa en secretaría. "
        "Si crees que otra persona la conoce, cámbiala de inmediato (Parte 7).",
        theme="azul",
    )

    add_heading_2(doc, "1.5 Si no puedes entrar")
    add_p(doc, "Lee el mensaje que aparece en rojo o en amarillo. Te dice qué pasa:")
    add_custom_table(
        doc,
        ["Mensaje", "Qué significa y qué hacer"],
        [
            ["Credenciales inválidas", "El usuario o la contraseña no coinciden. Revisa que pusiste ALU- delante del DNI, "
             "que no hay espacios y que la contraseña está bien escrita (las mayúsculas cuentan)."],
            ["Demasiados intentos fallidos", "Por seguridad, el campus se bloquea un rato después de varios intentos. "
             "Espera los minutos que indica el mensaje antes de volver a probar."],
            ["Cuenta desactivada", "Tu cuenta no está activa. Acércate a secretaría."],
            ["Tu sesión se cerró por inactividad", "Pasó mucho tiempo sin usar el campus. No es un error: vuelve a entrar."],
        ],
        [Inches(2.0), Inches(4.5)],
    )
    add_heading_3(doc, "El botón «No puedo entrar / Reportar un problema»")
    add_p(doc, "Está debajo del botón de iniciar sesión. Úsalo si olvidaste tu contraseña o no consigues entrar:")
    pasos(doc, [
        "Haz clic en «No puedo entrar / Reportar un problema».",
        "Escribe tu DNI (las 8 cifras) y un teléfono donde puedan llamarte.",
        "Cuenta qué te pasa, con tus palabras (al menos 20 caracteres). Por ejemplo, qué mensaje te sale.",
        "Presiona «Enviar solicitud». El colegio revisará tu caso y te llamará al teléfono que dejaste.",
    ])
    add_p(doc, "Ese aviso no cambia tu contraseña ni abre tu cuenta al instante: alguien del colegio lo atiende. "
               "Nunca escribas tu contraseña en ese formulario. También puedes acercarte a secretaría.", italic=True, color=COLOR_GRIS_OSCURO)

    add_heading_2(doc, "1.6 Cómo moverte por tu campus")
    add_p(doc, "A la izquierda tienes el menú guinda. «Alumno» y «Trámites» se abren y se cierran al hacer clic (fíjate en la flechita).")
    add_custom_table(
        doc,
        ["Menú", "Qué encuentras"],
        [
            ["Inicio", "Tu resumen: eventos, cursos, tareas pendientes y notificaciones."],
            ["Cursos", "Todas tus materias y el aula virtual de cada una."],
            ["Mensajería", "Chat con tus profesores y con psicología."],
            ["Alumno", "Conducta, Citas psicología, Notas, Asistencia, Horario y Matrícula."],
            ["Trámites", "Solicitud de trámite, Estado de cuenta y Manual de pagos."],
        ],
        [Inches(1.6), Inches(4.9)],
    )
    add_p(doc, "Arriba a la derecha están la campana de notificaciones y tu nombre. Al hacer clic en tu nombre se abre un menú con "
               "«Mis Datos», «Cambiar Contraseña» y «Cerrar Sesión».")
    add_p(doc, "En el celular el menú de la izquierda está escondido: se abre con el botón de las tres rayitas, arriba a la izquierda.")

    # ------------------------------------------------------- PARTE 2. INICIO
    add_heading_1(doc, "Parte 2. Tu página de Inicio")
    add_p(doc, "Es lo primero que ves al entrar. Es el resumen de tu día:")
    puntos(doc, [
        ("Próximos eventos ", "del colegio, para que no se te pase ninguna fecha."),
        ("Tus cursos ", "(o los que visitaste hace poco), para entrar directo a cualquiera."),
        ("Tareas pendientes. ", "Si no tienes ninguna, verás «¡Estás al día!»."),
        ("Notificaciones recientes: ", "los avisos nuevos del colegio."),
    ])
    add_p(doc, "Si te pierdes navegando, haz clic en «Inicio» y vuelves aquí.")
    add_captura_box(doc, 3, "Página de inicio del estudiante.")

    # ------------------------------------------------------- PARTE 3. CURSOS
    add_heading_1(doc, "Parte 3. Tus cursos y el aula virtual")

    add_heading_2(doc, "3.1 La lista de tus cursos")
    add_p(doc, "En «Cursos» ves todas tus materias del año, cada una con el nombre de su docente. "
               "Arriba puedes cambiar el año, por si quieres revisar uno anterior. Haz clic en un curso para entrar a su aula virtual.")
    add_p(doc, "Si dice «No tienes cursos asignados para este año», puede que tu matrícula aún no esté completa. Consúltalo en secretaría.")

    add_heading_2(doc, "3.2 Dentro de un curso")
    add_p(doc, "El aula virtual está ordenada por bimestres (I, II, III y IV). En cada bimestre hay dos bloques:")
    puntos(doc, [
        ("Contenido de clase: ", "los materiales que subió tu profesor (separatas, presentaciones, guías). Si tienen archivo, puedes descargarlo."),
        ("Tareas y exámenes: ", "las actividades de ese bimestre. Haz clic en una para abrirla."),
    ])
    add_captura_box(doc, 4, "Aula virtual de un curso con los bimestres, el contenido de clase y las tareas.")

    add_heading_2(doc, "3.3 Cómo entregar una tarea")
    pasos(doc, [
        "Entra al curso y busca la tarea en su bimestre.",
        "Haz clic en ella. Verás las instrucciones, la fecha de entrega y, si el profesor lo adjuntó, el botón «Descargar Documento Adjunto».",
        "En «Tu entrega», haz clic para elegir tu archivo. Se aceptan PDF, DOCX o imágenes, de hasta 10 MB.",
        "Presiona «Confirmar Entrega». Aparecerá el aviso «¡Tarea entregada con éxito!».",
    ])
    add_captura_box(doc, 5, "Ventana de una tarea con la zona para adjuntar el archivo.")
    add_callout(
        doc,
        "Revisa tu archivo antes de entregarlo",
        "Una vez entregada, la tarea queda marcada como «Tarea Entregada» y desde el campus ya no puedes cambiar el archivo. "
        "Asegúrate de subir el archivo correcto y completo. Si te equivocaste, escríbele a tu profesor por Mensajería.",
        theme="guinda",
    )

    add_heading_2(doc, "3.4 La nota y los comentarios de tu profesor")
    add_p(doc, "Si abres una tarea que ya entregaste, verás la «Retroalimentación del Docente» cuando tu profesor la escriba. "
               "La nota de cada evaluación y el promedio del bimestre los ves en «Alumno → Notas».")

    # ------------------------------------------------------- PARTE 4. ALUMNO
    add_heading_1(doc, "Parte 4. El menú «Alumno»")
    add_p(doc, "Aquí está todo tu registro: conducta, citas, notas, asistencia, horario y matrícula.")

    add_heading_2(doc, "4.1 Notas")
    pasos(doc, [
        "Entra a «Alumno → Notas».",
        "Si hace falta, cambia el año arriba a la derecha.",
        "En «Curso a consultar» elige un curso. Las notas se ven de un curso a la vez.",
    ])
    add_p(doc, "Verás tu promedio final del curso y, debajo, cada bimestre con su promedio y la nota de cada evaluación "
               "(algunas muestran su peso en porcentaje). Dos guiones (--) significan que todavía no hay nota.")
    add_captura_box(doc, 6, "Pantalla «Mis Notas» con un curso seleccionado.")
    add_callout(
        doc,
        "¿Tu promedio final se ve bajo?",
        "El promedio final es la suma de los cuatro bimestres dividida entre cuatro. Mientras el año avanza, "
        "los bimestres que aún no terminan cuentan como cero, así que el número se ve bajo. No es un error: "
        "se va corrigiendo solo cuando se cierra cada bimestre.",
        theme="azul",
    )

    add_heading_2(doc, "4.2 Asistencia")
    add_p(doc, "Muestra tu porcentaje de asistencia y tu historial día por día. Hay cuatro estados:")
    add_custom_table(
        doc,
        ["Estado", "Qué significa", "¿Afecta tu porcentaje?"],
        [
            ["Presente", "Llegaste a tiempo.", "Cuenta como asistencia."],
            ["Tardanza", "Llegaste después de la hora.", "Cuenta como asistencia, pero queda registrada."],
            ["Falta", "No viniste y no hay justificación.", "Sí, baja tu porcentaje."],
            ["Justificado", "No viniste, pero se presentó una justificación.", "No, no se cuenta."],
        ],
        [Inches(1.3), Inches(2.8), Inches(2.4)],
    )
    puntos(doc, [
        ("Filtro de mes: ", "arriba puedes elegir un mes o «Todo el año». El porcentaje y los contadores cambian según lo que elijas."),
        ("Contadores: ", "haz clic en Presente, Tardanza, Falta o Justificado para ver solo esos días."),
    ])
    add_p(doc, "Si ves un error (un día que viniste y sale falta), habla con tu auxiliar: es quien registra y corrige la asistencia.")
    add_captura_box(doc, 7, "Pantalla «Mi Asistencia» con el porcentaje, los contadores y el filtro de mes.")

    add_heading_2(doc, "4.3 Conducta")
    add_p(doc, "Tu conducta se mide en puntos y es tu nota de conducta de la libreta. "
               "Empiezas cada bimestre con 20 puntos, y cada falta registrada descuenta puntos según el Reglamento Interno. "
               "Al empezar un nuevo bimestre vuelves a 20.")
    add_custom_table(
        doc,
        ["Puntos", "Color", "Qué significa"],
        [
            ["20 a 15", "Verde · Buena conducta", "Todo en orden."],
            ["14 a 8", "Amarillo · En observación", "Has acumulado faltas. Cuida tu comportamiento."],
            ["7 a 0", "Rojo · Conducta crítica", "El colegio hará seguimiento de tu caso."],
        ],
        [Inches(1.2), Inches(2.3), Inches(3.0)],
    )
    add_p(doc, "Debajo ves tu última incidencia. Con «Ver historial completo» ves cada reporte con su fecha, la falta y la medida. "
               "Si crees que un reporte es un error, acude a tutoría o a psicología.")

    add_heading_2(doc, "4.4 Citas psicología")
    add_p(doc, "Aquí ves las citas que el área de psicología programó contigo: el motivo, la fecha, la hora y el lugar. "
               "La próxima cita sale destacada arriba. Con «Historial de citas» ves las que ya pasaron.")
    add_p(doc, "Desde esta pantalla no se piden citas. Si quieres hablar con un psicólogo, escríbele por Mensajería: "
               "si hace falta, te programará una cita y aparecerá aquí.")

    add_heading_2(doc, "4.5 Horario")
    add_p(doc, "Tu horario de clases de la semana, con la hora y el curso de cada bloque. "
               "Puedes descargarlo en PDF para imprimirlo o guardarlo en el celular.")

    add_heading_2(doc, "4.6 Matrícula")
    add_p(doc, "Muestra tu situación en el año escolar: año, estado, grado y sección. Si algo no corresponde, avisa en secretaría cuanto antes.")
    add_p(doc, "Cuando el colegio abre las inscripciones, desde aquí también puedes:")
    puntos(doc, [
        ("Solicitar la renovación de matrícula ", "para el año siguiente y ver la «Respuesta del Colegio»."),
        ("Inscribirte al ciclo de verano, ", "eligiendo cursos fijos, talleres o ambos. Se genera un pago que debes cancelar para ser admitido."),
    ])
    add_p(doc, "Si te sale el aviso «Tienes pagos pendientes», primero tendrás que ponerte al día.")

    # ----------------------------------------------------- PARTE 5. TRÁMITES
    add_heading_1(doc, "Parte 5. Trámites y pagos")

    add_heading_2(doc, "5.1 Estado de cuenta")
    add_p(doc, "En «Trámites → Estado de cuenta» ves de un vistazo:")
    puntos(doc, [
        ("Total Pendiente: ", "lo que falta pagar."),
        ("Pagos Realizados: ", "lo que ya está pagado y registrado."),
        ("Pensiones pendientes: ", "cada cuota con su fecha de vencimiento."),
        ("Historial de pagos: ", "cada pago con su fecha, concepto, número de operación e importe. «Ver todo el historial» abre la lista completa."),
    ])
    add_captura_box(doc, 8, "Pantalla «Trámites y Pagos» con el total pendiente.")

    add_heading_2(doc, "5.2 Cómo se pagan las pensiones")
    add_p(doc, "Los pagos NO se hacen dentro del campus. Se hacen en el Banco de Crédito del Perú (BCP):")
    puntos(doc, [
        ("En un agente BCP: ", "indica el servicio Montehermozo SAC, da el DNI del alumno y el monto."),
        ("En la app Banca Móvil BCP: ", "en «Pagos de servicios» busca Montehermozo SAC e ingresa el DNI del alumno."),
    ])
    add_p(doc, "El pago no aparece en el campus al instante: el colegio lo registra cuando recibe el reporte del banco. "
               "Si pasan varios días y no aparece, guarda tu comprobante y acércate a secretaría.")
    add_p(doc, "En «Trámites → Manual de pagos» tienes estas instrucciones siempre a mano, junto con el horario de secretaría.")

    add_heading_2(doc, "5.3 Cómo pedir un trámite")
    pasos(doc, [
        "Entra a «Trámites → Solicitud de trámite».",
        "Elige el trámite de la lista. Al lado de cada uno ves su costo o la palabra GRATUITO.",
        "Lee el recuadro azul: te dice los requisitos y, si tiene costo, cuántos días tienes para pagarlo.",
        "Adjunta tu documento de sustento si hace falta (PDF, JPG o PNG).",
        "Si quieres, escribe un comentario explicando para qué lo necesitas.",
        "Presiona «Enviar Solicitud». Tu solicitud aparecerá en la lista con su estado.",
    ])
    add_captura_box(doc, 9, "Pantalla de solicitud de trámite con el recuadro de información.")
    add_callout(
        doc,
        "Dos reglas que conviene saber",
        "Los trámites GRATUITOS (certificados, justificaciones y similares) no se pueden enviar sin adjuntar un documento de sustento: "
        "tenlo escaneado o fotografiado antes de empezar. Y los que dicen «requiere estar al día» no se pueden pedir si tienes cuotas vencidas: "
        "sale un aviso amarillo con cuántas cuotas debes, y el botón cambia a «Pensiones pendientes» y no funciona. "
        "Cuando el colegio registre tu pago, el botón vuelve a activarse.",
        theme="guinda",
    )

    add_heading_2(doc, "5.4 Los estados de tu solicitud")
    add_custom_table(
        doc,
        ["Estado", "Qué significa"],
        [
            ["Pendiente Pago", "El trámite tiene costo y todavía no se registra el pago. Verás la fecha límite."],
            ["En Revisión", "El pago ya está registrado y el colegio está revisando tu solicitud."],
            ["Aprobado", "Tu trámite fue aceptado."],
            ["Rechazado", "No fue aceptado. Lee la «Respuesta del Colegio» para saber por qué."],
        ],
        [Inches(1.8), Inches(4.7)],
    )
    add_p(doc, "Cuando el colegio te responde, la respuesta aparece dentro de tu solicitud. No hace falta llamar a preguntar.")

    # --------------------------------------------------- PARTE 6. MENSAJERÍA
    add_heading_1(doc, "Parte 6. Mensajería")
    add_p(doc, "Desde «Mensajería» puedes escribir a:")
    puntos(doc, [
        ("Tus profesores, ", "pero solo los que te dictan clase este año."),
        ("El área de psicología, ", "siempre."),
    ])
    add_p(doc, "No puedes iniciar conversaciones con otros alumnos. Si alguien de la administración del colegio te escribe, sí puedes responderle.")
    pasos(doc, [
        "Entra a «Mensajería».",
        "Busca a la persona por su nombre y abre la conversación.",
        "Escribe tu mensaje abajo y envíalo.",
    ])
    add_callout(
        doc,
        "Es un canal del colegio",
        "Todo lo que escribes queda registrado. Escribe con respeto, como le hablarías a un profesor en persona, "
        "y ten paciencia con la respuesta: los profesores pasan buena parte del día en clase.",
        theme="azul",
    )

    # ---------------------------------------------- PARTE 7. PERFIL Y CUENTA
    add_heading_1(doc, "Parte 7. Tus datos, tu contraseña y tus notificaciones")

    add_heading_2(doc, "7.1 Mis Datos")
    add_p(doc, "Haz clic en tu nombre (arriba a la derecha) y elige «Mis Datos». Tiene tres pestañas:")
    puntos(doc, [
        ("Datos personales: ", "puedes actualizar tu teléfono, tu correo y tu dirección. Tu nombre y tu DNI no se cambian desde aquí: si están mal, avisa en secretaría."),
        ("Datos médicos: ", "alergias o condiciones de salud que el colegio debe conocer."),
        ("Familiares: ", "tus familiares registrados. Puedes agregar uno nuevo."),
    ])

    add_heading_2(doc, "7.2 Cambiar tu contraseña cuando quieras")
    add_p(doc, "Además del cambio obligatorio de la primera vez, puedes cambiar tu contraseña siempre que quieras. "
               "Hazlo si crees que alguien la conoce, si la escribiste en una computadora ajena o si hace mucho que no la cambias.")
    pasos(doc, [
        "Haz clic en tu nombre y elige «Cambiar Contraseña».",
        "Escribe tu contraseña actual.",
        "Escribe la nueva (al menos 8 caracteres, que no sea tu DNI) y repítela. Busca que la barra diga Fuerte.",
        "Presiona «Actualizar Contraseña».",
    ])

    add_heading_2(doc, "7.3 Notificaciones")
    add_p(doc, "La campana de arriba a la derecha muestra los avisos del colegio. El número sobre la campana indica cuántos no has visto.")

    add_heading_2(doc, "7.4 Cierra tu sesión")
    add_p(doc, "Si usas una computadora que no es tuya (de la biblioteca, de un familiar, de una cabina), al terminar haz clic en tu nombre "
               "y elige «Cerrar Sesión». Si no lo haces, la siguiente persona podría usar tu cuenta.")

    # ------------------------------------------- PARTE 8. SEGURIDAD (RESUMEN)
    add_heading_1(doc, "Parte 8. Cuida tu cuenta")
    add_custom_table(
        doc,
        ["Revisa", "Por qué"],
        [
            ["Cambiaste la contraseña inicial (tu DNI).", "Tu DNI lo conoce mucha gente."],
            ["Tu contraseña tiene 10 caracteres o más, con letras y números.", "Así es muy difícil de adivinar."],
            ["No se la has dado a nadie.", "Nadie del colegio te la pedirá nunca."],
            ["No la usas en otras páginas o juegos.", "Si otra página se filtra, tu cuenta sigue a salvo."],
            ["Cierras sesión en computadoras ajenas.", "Para que nadie entre con tu cuenta abierta."],
            ["El correo registrado es de tus padres o apoderado.", "Ahí llegan los avisos del colegio."],
        ],
        [Inches(3.4), Inches(3.1)],
    )

    # --------------------------------------------------- PARTE 9. PREGUNTAS
    add_heading_1(doc, "Parte 9. Preguntas frecuentes")
    preguntas = [
        ("Escribo mi DNI y me sale «Credenciales inválidas».",
         "Tu usuario no es solo el DNI: escribe ALU- delante, por ejemplo ALU-73715514. La primera vez, la contraseña sí es solo tu DNI."),
        ("Olvidé mi contraseña, ¿qué hago?",
         "En la pantalla de acceso usa «No puedo entrar / Reportar un problema» o acércate a secretaría para que te la restablezcan."),
        ("El campus no me deja usar mi DNI como contraseña nueva.",
         "Es a propósito: la contraseña nueva no puede ser tu DNI. Elige otra de al menos 8 caracteres."),
        ("Me sale «Demasiados intentos fallidos».",
         "Espera los minutos que indica el mensaje y vuelve a intentarlo con calma, revisando el usuario y la contraseña."),
        ("No me aparece ninguna nota en un curso.",
         "Puede que tu profesor aún no haya registrado las notas de ese bimestre. Si pasa un tiempo y sigue vacío, pregúntale por Mensajería."),
        ("Mi promedio final sale muy bajo y mis notas son buenas.",
         "El promedio final divide entre cuatro bimestres, y los que aún no terminan cuentan como cero. Se corrige solo durante el año."),
        ("Falté con justificación pero aparece como falta.",
         "Presenta la justificación a tu auxiliar para que cambie el registro a Justificado. Así no afecta tu porcentaje."),
        ("Entregué el archivo equivocado en una tarea.",
         "Desde el campus no se puede cambiar. Escríbele a tu profesor por Mensajería lo antes posible."),
        ("No puedo enviar mi solicitud de trámite.",
         "Si es gratuito, revisa que adjuntaste el sustento. Si el botón dice «Pensiones pendientes», primero ponte al día con los pagos."),
        ("Ya pagué pero sigue apareciendo como pendiente.",
         "El pago del banco tarda en registrarse. Si pasan varios días, acércate a secretaría con tu comprobante."),
        ("¿Puedo pagar la pensión desde el campus?",
         "No. Se paga en el BCP, en agente o por la app, con el servicio Montehermozo SAC y el DNI del alumno."),
        ("No puedo escribirle a un profesor.",
         "Solo puedes escribir a los profesores que te dictan clase este año. Para otro, hazlo a través de tu tutor o de secretaría."),
        ("¿Puedo usar el campus desde el celular?",
         "Sí, se adapta a la pantalla. El menú se abre con las tres rayitas de arriba a la izquierda."),
    ]
    for pregunta, respuesta in preguntas:
        add_heading_3(doc, pregunta)
        add_p(doc, respuesta)

    add_heading_1(doc, "¿Necesitas ayuda?")
    add_p(doc, "Si algo no funciona como dice este manual, avisa a tu tutor, a tu auxiliar o en secretaría. "
               "Atención en secretaría: de 7:30 a. m. a 1:00 p. m. y de 3:00 p. m. a 6:30 p. m.")

    doc.save(DOC_PATH)
    print(f"Manual del Estudiante guardado en: {DOC_PATH}")


if __name__ == "__main__":
    construir_manual()
